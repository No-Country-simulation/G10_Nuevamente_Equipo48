"""Validación determinista de consistencia del corpus NuevaMente Fase 05."""
from __future__ import annotations
import csv, json, sys
from pathlib import Path
from collections import Counter
try:
    import openpyxl
except ImportError:
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]
INV = ROOT / "inventario" / "inventario_corpus.csv"
EVID = ROOT / "reportes" / "registro_evidencias_fuentes.csv"
JSON = ROOT / "ejemplos" / "ejemplo_metadatos_documento.json"
XLSX = ROOT / "matriz" / "Matriz_Corpus_NuevaMente_Consolidada_v2_Poblacion.xlsx"
RESULTS = {"APTO", "APTO_CON_REVISIÓN", "NO_APTO"}
EXPECTED_COUNTS = {"APTO": 18, "APTO_CON_REVISIÓN": 2, "NO_APTO": 2}

def load_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def main() -> int:
    errors, warnings = [], []
    inv, evid = [], []
    if not INV.exists(): errors.append(f"No existe {INV}")
    else:
        inv = load_csv(INV)
        if len(inv) != 22: errors.append(f"Inventario: se esperaban 22 registros y hay {len(inv)}")
        required = {"identificador_documento", "titulo", "fuente", "direccion_fuente", "idioma", "formato", "dominio", "licencia", "estado_seleccion", "resultado_evaluacion"}
        if inv and required - set(inv[0]): errors.append(f"Faltan columnas: {sorted(required-set(inv[0]))}")
        ids = [r.get("identificador_documento", "").strip() for r in inv]
        if any(not x for x in ids): errors.append("Hay identificadores vacíos en el inventario")
        if len(ids) != len(set(ids)): errors.append("Hay identificadores duplicados en el inventario")
        for r in inv:
            if r.get("resultado_evaluacion") not in RESULTS: errors.append(f"{r.get('identificador_documento')}: resultado inválido")
            if r.get("estado_seleccion") != "EN_EVALUACION": errors.append(f"{r.get('identificador_documento')}: estado de selección inesperado")
            if r.get("observacion_adaptacion_nna", "").strip().endswith(("...", "…")): errors.append(f"{r.get('identificador_documento')}: observación NNA truncada")
        counts = dict(Counter(r.get("resultado_evaluacion", "") for r in inv))
        if counts != EXPECTED_COUNTS: errors.append(f"Distribución inesperada: {counts}")
    if not EVID.exists(): errors.append(f"No existe {EVID}")
    else:
        evid = load_csv(EVID)
        if len(evid) != 22: errors.append(f"Evidencias: se esperaban 22 registros y hay {len(evid)}")
        imap = {r.get("identificador_documento"): r for r in inv}
        emap = {r.get("ID"): r for r in evid}
        if set(imap) != set(emap): errors.append("Los IDs de inventario y evidencias no coinciden")
        for rid in set(imap) & set(emap):
            if imap[rid].get("resultado_evaluacion") != emap[rid].get("Resultado_auditoria"):
                errors.append(f"{rid}: resultado distinto entre inventario y evidencias")
            if imap[rid].get("direccion_fuente") != emap[rid].get("URL_fuente_principal"):
                errors.append(f"{rid}: URL distinta entre inventario y evidencias")
    if not JSON.exists(): errors.append(f"No existe {JSON}")
    else:
        try:
            data = json.loads(JSON.read_text(encoding="utf-8"))
            if not isinstance(data, dict): errors.append("El JSON de ejemplo debe tener un objeto raíz")
        except json.JSONDecodeError as exc: errors.append(f"JSON inválido: {exc}")
    if not XLSX.exists(): errors.append(f"No existe {XLSX}")
    elif openpyxl is None: errors.append("No está disponible openpyxl para validar la matriz")
    else:
        wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
        required_sheets = {"Matriz maestra", "Finalistas propuestos", "Auditoría integral"}
        if not required_sheets.issubset(set(wb.sheetnames)): errors.append(f"Faltan hojas: {sorted(required_sheets-set(wb.sheetnames))}")
        else:
            def sheet_records(name):
                ws = wb[name]
                values = list(ws.values)
                headers = values[0]
                return [dict(zip(headers, row)) for row in values[1:] if row and row[0]]
            matrix = sheet_records("Matriz maestra")
            audit = sheet_records("Auditoría integral")
            imap = {r.get("identificador_documento"): r for r in inv}
            for label, records, id_key, result_key in [
                ("Matriz maestra", matrix, "ID", "Resultado"),
                ("Auditoría integral", audit, "ID", "Resultado auditado")]:
                rmap = {r.get(id_key): r for r in records}
                if set(rmap) != set(imap): errors.append(f"IDs de inventario y {label} no coinciden")
                for rid in set(rmap) & set(imap):
                    if rmap[rid].get(result_key) != imap[rid].get("resultado_evaluacion"):
                        errors.append(f"{rid}: resultado distinto en {label}")
            finalists = sheet_records("Finalistas propuestos")
            if len(finalists) != 8: warnings.append(f"La hoja de finalistas contiene {len(finalists)} propuestas (referencia histórica: 8)")
            for r in finalists:
                rid = r.get("ID")
                if rid not in imap: errors.append(f"Finalista {rid} no existe en inventario")
                elif imap[rid].get("resultado_evaluacion") != "APTO": errors.append(f"Finalista {rid} no está APTO")
        wb.close()
    print(f"Inventario: {len(inv)} registros")
    print(f"Evidencias: {len(evid)} registros")
    print("Matriz maestra y auditoría integral: verificadas" if XLSX.exists() and openpyxl else "Matriz: no verificada")
    if inv: print(f"Distribución: {dict(Counter(r.get('resultado_evaluacion','') for r in inv))}")
    print(f"ERRORES: {len(errors)}")
    for e in errors: print(f"  ERROR: {e}")
    print(f"ADVERTENCIAS: {len(warnings)}")
    for w in warnings: print(f"  AVISO: {w}")
    return 1 if errors else 0

if __name__ == "__main__":
    raise SystemExit(main())
